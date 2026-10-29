"""
reports/workbook.py — the writer behind every export.

  * a genuine .xlsx (zip container openpyxl reads back), not renamed CSV
  * frozen, bold, filterable header row; sized columns; usable sheet names
  * typed cells: dates, IST datetimes, rupee amounts, Yes/No, real blanks
  * customer-typed text can never become a formula, and control characters
    are stripped instead of breaking the file
"""
import datetime
import io
import zipfile

from django.test import SimpleTestCase
from openpyxl import load_workbook

from bookings.slots import IST

from .workbook import (
    DATE, DATETIME, INT, MONEY, Cell, Sheet, build, col, excel_value, sheet_title,
    xlsx_response,
)

COLUMNS = [col('Name', 20), col('Amount', 12, MONEY), col('Day', 12, DATE),
           col('When', 18, DATETIME), col('Count', 8, INT), col('Flag', 6)]


def load(content):
    return load_workbook(io.BytesIO(content))


class WorkbookTests(SimpleTestCase):

    def test_is_a_real_xlsx_workbook(self):
        content = build([Sheet('Bookings', COLUMNS, [['Asha', 100, None, None, 1, True]])])
        self.assertTrue(content.startswith(b'PK'))
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
        self.assertIn('xl/workbook.xml', names)
        self.assertIn('[Content_Types].xml', names)
        self.assertEqual(load(content).sheetnames, ['Bookings'])

    def test_header_is_bold_frozen_filtered_and_sized(self):
        rows = [['A', 1, None, None, 1, False], ['B', 2, None, None, 2, True]]
        ws = load(build([Sheet('Data', COLUMNS, rows)]))['Data']
        self.assertEqual([c.value for c in ws[1]], [h for h, _w, _f in COLUMNS])
        self.assertTrue(ws['A1'].font.bold)
        self.assertEqual(ws['A1'].fill.fgColor.rgb[-6:], '1F3A5F')
        self.assertEqual(ws.freeze_panes, 'A2')
        self.assertEqual(ws.auto_filter.ref, 'A1:F3')
        self.assertEqual(ws.column_dimensions['A'].width, 20)
        self.assertEqual(ws.column_dimensions['D'].width, 18)

    def test_cells_are_typed(self):
        when = datetime.datetime(2026, 10, 24, 6, 30, tzinfo=datetime.timezone.utc)
        day = datetime.date(2026, 10, 25)
        ws = load(build([Sheet('Data', COLUMNS, [['Asha', 1220, day, when, 3, True]])]))['Data']
        name, amount, date_cell, when_cell, count, flag = ws[2]
        self.assertEqual(amount.value, 1220)
        self.assertEqual(amount.number_format, MONEY)
        self.assertEqual(date_cell.value.date(), day)
        self.assertEqual(date_cell.number_format, DATE)
        # 06:30 UTC is 12:00 in India; Excel holds the IST wall-clock time.
        self.assertEqual(when_cell.value, datetime.datetime(2026, 10, 24, 12, 0))
        self.assertEqual(when_cell.number_format, DATETIME)
        self.assertEqual(count.value, 3)
        self.assertEqual(flag.value, 'Yes')

    def test_blank_values_are_empty_cells(self):
        ws = load(build([Sheet('Data', COLUMNS, [['', None, None, None, None, '']])]))['Data']
        self.assertTrue(all(c.value is None for c in ws[2]))

    def test_text_that_looks_like_a_formula_stays_text(self):
        evil = '=HYPERLINK("http://evil.example","click")'
        ws = load(build([Sheet('Data', [col('Name', 20)], [[evil], ['+91 98765'], ['@sum']])]))['Data']
        self.assertEqual(ws['A2'].value, evil)
        self.assertEqual(ws['A2'].data_type, 's')
        self.assertEqual(ws['A3'].data_type, 's')
        self.assertEqual(ws['A4'].data_type, 's')
        with zipfile.ZipFile(io.BytesIO(build([Sheet('D', [col('N', 5)], [[evil]])]))) as archive:
            sheet_xml = archive.read('xl/worksheets/sheet1.xml').decode()
        self.assertNotIn('<f>', sheet_xml)

    def test_control_characters_are_stripped_not_fatal(self):
        ws = load(build([Sheet('Data', [col('Note', 20)], [['Cake\x00 at \x078']])]))['Data']
        self.assertEqual(ws['A2'].value, 'Cake at 8')

    def test_empty_sheet_keeps_its_header(self):
        ws = load(build([Sheet('Empty', COLUMNS, [])]))['Empty']
        self.assertEqual(ws.max_row, 1)
        self.assertEqual([c.value for c in ws[1]][0], 'Name')
        self.assertEqual(ws.auto_filter.ref, 'A1:F2')

    def test_several_sheets_in_order(self):
        content = build([Sheet('Summary', COLUMNS, [], filters=False),
                         Sheet('Bookings', COLUMNS, []), Sheet('Refunds', COLUMNS, [])])
        wb = load(content)
        self.assertEqual(wb.sheetnames, ['Summary', 'Bookings', 'Refunds'])
        self.assertIsNone(wb['Summary'].auto_filter.ref)

    def test_cell_overrides_format_and_bold(self):
        ws = load(build([Sheet('S', [col('Metric', 20), col('Value', 12)], [
            ['Gross', Cell(5000, MONEY, bold=True)], ['Bookings', Cell(4, INT)],
        ], filters=False)]))['S']
        self.assertEqual(ws['B2'].number_format, MONEY)
        self.assertTrue(ws['B2'].font.bold)
        self.assertEqual(ws['B3'].number_format, INT)

    def test_sheet_names_are_made_valid(self):
        self.assertEqual(sheet_title('Payments [Oct/Nov]?'), 'Payments  Oct Nov')
        self.assertEqual(len(sheet_title('x' * 50)), 31)
        self.assertEqual(sheet_title(''), 'Sheet')

    def test_excel_value(self):
        naive = datetime.datetime(2026, 1, 1, 9, 0)
        self.assertEqual(excel_value(naive), naive)
        aware = datetime.datetime(2026, 1, 1, 9, 0, tzinfo=IST)
        self.assertEqual(excel_value(aware), naive)
        self.assertEqual(excel_value(False), 'No')
        self.assertIsNone(excel_value(''))
        self.assertEqual(excel_value(0), 0)

    def test_download_response_is_an_uncached_attachment(self):
        response = xlsx_response(b'PK..', 'admin_users_2026-10-24.xlsx')
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertEqual(response['Content-Disposition'],
                         'attachment; filename="admin_users_2026-10-24.xlsx"')
        self.assertEqual(response['Cache-Control'], 'no-store')
