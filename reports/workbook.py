"""
The ONE .xlsx writer every export uses (openpyxl, write-only mode).

A workbook is a list of Sheets; a Sheet is a title, its columns and its rows.
Every sheet comes out the same way: a bold, filled header row that is frozen
and filterable, columns sized for their content, and typed cells — real dates
and datetimes, rupee amounts as numbers with a ₹ format — so the file sorts,
filters and sums correctly in Excel instead of holding text that looks like
numbers.

Columns are plain (header, width, number_format) tuples — the shape
adminpanel/exports.py has always used for the Bookings export — with the
formats below as constants. A single cell can override its column's format
or be bold through Cell (used by the Summary sheets, whose value column mixes
counts and amounts).

SAFETY: openpyxl stores any string that starts with "=" as a live formula.
Names, notes and search terms come from customers and vendors, so a booking
made by "=HYPERLINK(…)" would otherwise become a formula in the admin's
spreadsheet. Every string is forced to be text, and characters XML cannot
carry (control codes pasted into a note) are stripped instead of failing the
whole export.
"""
import datetime
import io
import re

from bookings.slots import IST

# Number formats.
MONEY = '"₹"#,##0'
INT = '0'
DATE = 'dd-mmm-yyyy'
DATETIME = 'dd-mmm-yyyy hh:mm'
PERCENT = '0%'

CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

HEADER_FILL = '1F3A5F'
HEADER_FONT = 'FFFFFF'

# Excel's own limits on a sheet name.
_SHEET_NAME_BAD = re.compile(r'[\[\]:*?/\\]')
_SHEET_NAME_MAX = 31


def col(header, width, number_format=None):
    """One column: header text, width in characters, optional number format."""
    return (header, width, number_format)


class Cell:
    """A value that overrides its column's number format, or is bold."""

    __slots__ = ('value', 'number_format', 'bold')

    def __init__(self, value, number_format=None, bold=False):
        self.value = value
        self.number_format = number_format
        self.bold = bold


class Sheet:
    """One worksheet. `rows` is any iterable of lists, one value per column.

    `filters=False` leaves off the header's filter buttons (a two-column
    summary has nothing to filter)."""

    def __init__(self, title, columns, rows, *, filters=True):
        self.title = sheet_title(title)
        self.columns = list(columns)
        self.rows = rows
        self.filters = filters


def sheet_title(title):
    """A name Excel accepts: no []:*?/\\ and at most 31 characters."""
    clean = _SHEET_NAME_BAD.sub(' ', str(title or 'Sheet')).strip() or 'Sheet'
    return clean[:_SHEET_NAME_MAX]


def excel_value(value):
    """A Python value -> what the cell should hold.

    Aware datetimes become IST wall-clock time (Excel has no time zones),
    booleans read Yes/No, and an empty string is a truly blank cell."""
    if isinstance(value, datetime.datetime):
        if value.tzinfo is not None:
            value = value.astimezone(IST).replace(tzinfo=None)
        return value
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    if value == '':
        return None
    if isinstance(value, str):
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        value = ILLEGAL_CHARACTERS_RE.sub('', value)
        return value or None
    return value


def build(sheets):
    """The .xlsx file as bytes."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    wb.properties.creator = 'TheBookMyVenues'
    header_font = Font(bold=True, color=HEADER_FONT)
    header_fill = PatternFill('solid', fgColor=HEADER_FILL)
    bold_font = Font(bold=True)

    for sheet in sheets:
        ws = wb.create_sheet(sheet.title)
        ws.freeze_panes = 'A2'
        for index, (_header, width, _fmt) in enumerate(sheet.columns, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width

        header = []
        for title, _width, _fmt in sheet.columns:
            cell = WriteOnlyCell(ws, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical='center')
            header.append(cell)
        ws.append(header)

        formats = [fmt for _h, _w, fmt in sheet.columns]
        count = 0
        for row in sheet.rows:
            out = []
            for value, number_format in zip(row, formats):
                bold = False
                if isinstance(value, Cell):
                    number_format = value.number_format or number_format
                    bold = value.bold
                    value = value.value
                value = excel_value(value)
                text = isinstance(value, str)
                if bold or text or (number_format and value is not None):
                    cell = WriteOnlyCell(ws, value=value)
                    if text:
                        cell.data_type = 's'   # never a formula
                    elif number_format and value is not None:
                        cell.number_format = number_format
                    if bold:
                        cell.font = bold_font
                    out.append(cell)
                else:
                    out.append(value)
            ws.append(out)
            count += 1

        if sheet.filters:
            last_col = get_column_letter(len(sheet.columns))
            ws.auto_filter.ref = f'A1:{last_col}{max(count, 1) + 1}'

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def xlsx_response(content, filename):
    """The download response: the right type, an attachment name, and never
    cached anywhere — every export carries customer or business data."""
    from django.http import HttpResponse

    response = HttpResponse(content, content_type=CONTENT_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Cache-Control'] = 'no-store'
    return response
