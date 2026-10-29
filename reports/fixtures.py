"""
Test data shared by the export tests: two vendors, their venues (with a unit
listing and a deleted one), venue bookings in every payment state, events
with ticket orders in every state, payouts and audit entries.

Not a test module (no `test` prefix) — imported by reports/test_*.py.
"""
import datetime
import io
import uuid
import zipfile

from django.utils import timezone
from openpyxl import load_workbook

from adminpanel.models import AuditEntry, Payout
from bookings.models import Booking
from bookings.slots import today_ist
from bookings.views import PENDING_HOLD_MINUTES
from events.models import Event, EventOrder
from events.tests import make_admin, make_customer, make_event, make_vendor
from venues.models import Listing

OLD = PENDING_HOLD_MINUTES + 5


def read_sheets(response):
    """{sheet title: [ {header: value} ]} from an .xlsx response."""
    wb = load_workbook(io.BytesIO(response.content))
    out = {}
    for ws in wb.worksheets:
        values = list(ws.iter_rows(values_only=True))
        headers = list(values[0]) if values else []
        out[ws.title] = [dict(zip(headers, row)) for row in values[1:]]
    return out


def headers(response, sheet):
    wb = load_workbook(io.BytesIO(response.content))
    return [c.value for c in wb[sheet][1]]


def raw_xml(response):
    """Every XML part of the workbook, as one string — to prove a value is
    NOT in the file anywhere."""
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        return ''.join(
            archive.read(name).decode('utf-8', 'replace')
            for name in archive.namelist() if name.endswith('.xml')
        )


class World:
    """Builds the shared data set. Attributes name every object a test
    asserts on."""

    def __init__(self):
        self.today = today_ist()
        self.admin = make_admin()
        self.vendor = make_vendor('9000000001', 'Vendor One')
        self.other_vendor = make_vendor('9000000002', 'Vendor Two')
        self.asha = make_customer('9100000001', 'Asha')
        self.ravi = make_customer('9100000002', 'Ravi')
        self.ravi.is_active = False
        self.ravi.save(update_fields=['is_active'])
        self.asha.set_password('secret-password-123')
        self.asha.save()

        L = self.listing
        self.hall = L(self.vendor, 'Grand Hall', price=1200, city='Warangal')
        self.screen = L(self.vendor, 'Grand Hall — Screen 2', unit_of=self.hall)
        self.turf = L(self.vendor, 'City Turf', status=Listing.Status.PAUSED, price=800)
        self.old = L(self.vendor, 'Old Place', status=Listing.Status.DELETED)
        self.arena = L(self.other_vendor, 'Rival Arena', price=2000)
        self.pending = L(self.vendor, 'New Banquet', status=Listing.Status.PENDING,
                         city='Hanamkonda', category='Party hall')

        B = self.booking
        self.b_paid = B(self.hall, 'confirmed', days=2, pay='pay_1', method='upi')
        self.b_done = B(self.screen, 'confirmed', days=-3, pay='pay_2', method='card',
                        amount=1020, unit_label='Screen 2')
        self.b_walkin = B(self.hall, 'confirmed', days=1, amount=500, fee=0,
                          method='walk-in', walk_in=True, customer='Walk-in',
                          phone='', user=None)
        self.b_part = B(self.hall, 'confirmed', days=3, pay='pay_3', amount=2020,
                        pay_now=520, at_venue=1500)
        self.b_promo = B(self.hall, 'confirmed', days=6, pay='pay_4', amount=900,
                         discount=100, offer={'code': 'BMV100', 'source': 'platform'})
        self.b_refund_pending = B(self.turf, 'refund_pending', days=5, pay='pay_5', amount=820)
        self.b_refunded = B(self.turf, 'refunded', days=-1, pay='pay_6', amount=620,
                            refund_amount=300, refund_reason='Rain')
        self.b_refunded_full = B(self.hall, 'refunded', days=-2, pay='pay_7', amount=720)
        self.b_hold = B(self.hall, 'payment_pending', days=4)
        self.b_cancelled = B(self.hall, 'cancelled', days=4)
        self.b_cancelled_paid = B(self.hall, 'cancelled', days=4, pay='pay_8', amount=410)
        self.b_failed = B(self.hall, 'failed', days=4)
        self.b_removed = B(self.hall, 'removed', days=4)
        self.b_other = B(self.arena, 'confirmed', days=2, pay='pay_9', amount=5000,
                         user=self.ravi, customer='Ravi', phone='9100000002')

        # Events: an admin event, the vendor's own (pending) event, a
        # published admin event at the vendor's venue, and the other vendor's
        # event at the other vendor's venue.
        self.garba = make_event(name='Garba Night', days_ahead=(1, 2), capacity=100)
        self.own_event = make_event(self.vendor, name='Vendor One Fest',
                                    status=Event.Status.PENDING, days_ahead=(5,))
        self.hall_event = make_event(name='Hall Night', days_ahead=(3,), listing=self.hall)
        self.rival_event = make_event(self.other_vendor, name='Rival Fest',
                                      days_ahead=(4,), listing=self.arena)

        O = self.order
        self.o_paid = O(self.garba, 'confirmed', pay='pay_e1')
        self.o_paid2 = O(self.garba, 'confirmed', pay='pay_e2', adult=1, child=0)
        self.o_refund_pending = O(self.garba, 'refund_pending', pay='pay_e3')
        self.o_refunded = O(self.garba, 'refunded', pay='pay_e4')
        self.o_cancelled_paid = O(self.garba, 'cancelled', pay='pay_e5')
        self.o_hold = O(self.garba, 'payment_pending')
        self.o_failed = O(self.garba, 'failed')
        self.o_own = O(self.own_event, 'confirmed', pay='pay_e6')
        self.o_hall = O(self.hall_event, 'confirmed', pay='pay_e7')
        self.o_hall_hold = O(self.hall_event, 'payment_pending')
        self.o_rival = O(self.rival_event, 'confirmed', pay='pay_e8',
                         customer='Rival Buyer', phone='9100000099')

        monday = self.today - datetime.timedelta(days=self.today.weekday() + 14)
        self.payout_pending = Payout.objects.create(
            vendor='Vendor One', vendor_user=self.vendor, period='wk',
            period_start=monday, period_end=monday + datetime.timedelta(days=6),
            gross=4321, status=Payout.Status.PENDING,
        )
        self.payout_done = Payout.objects.create(
            vendor='Vendor Two', vendor_user=self.other_vendor, period='wk',
            period_start=monday - datetime.timedelta(days=7),
            period_end=monday - datetime.timedelta(days=1),
            gross=999, status=Payout.Status.COMPLETED,
        )

    # ---- factories --------------------------------------------------

    def listing(self, vendor, name, *, status=Listing.Status.LIVE, unit_of=None,
                price=1000, city='', category='Party hall'):
        detail = {'capacity': 50}
        if unit_of is not None:
            detail['unitOf'] = str(unit_of.id)
        listing = Listing.objects.create(
            id=uuid.uuid4(), vendor=vendor,
            slug=name.lower().replace(' ', '-').replace('—', 'x') + uuid.uuid4().hex[:6],
            record={'name': name, 'price': price, 'city': city, 'district': city,
                    'category': category, 'location': f'{name} Road', 'detail': detail},
            name=name, category=category, locality='Hanamkonda', pincode='506001',
            status=status,
        )
        if status == Listing.Status.DELETED:
            Listing.objects.filter(pk=listing.pk).update(deleted_at=timezone.now())
        return listing

    def booking(self, listing, status, *, days, pay='', method='online', amount=1220,
                fee=20, customer='Asha', phone='9100000001', user='asha', walk_in=False,
                pay_now=None, at_venue=0, discount=0, offer=None, refund_amount=None,
                refund_reason='', unit_label='', age=0):
        booking = Booking.objects.create(
            listing=listing, user=self.asha if user == 'asha' else user,
            venue_name=listing.name, category=listing.category,
            date=self.today + datetime.timedelta(days=days),
            slots=['19:00 – 20:00'], amount=amount, fee=fee, customer_name=customer,
            phone=phone, status=status, razorpay_payment_id=pay,
            razorpay_order_id=f'order_{pay}' if pay else '', method=method,
            walk_in=walk_in, pay_now=pay_now, at_venue=at_venue,
            discount_amount=discount, offer=offer, refund_amount=refund_amount,
            refund_reason=refund_reason, unit_label=unit_label,
            addons=[{'name': 'Cake', 'qty': 1, 'price': 500}],
        )
        if age:
            Booking.objects.filter(pk=booking.pk).update(
                created_at=timezone.now() - datetime.timedelta(minutes=age)
            )
        return booking

    def order(self, event, status, *, pay='', adult=2, child=1, customer='Asha',
              phone='9100000001', age=0):
        event_date = event.dates.first()
        items = []
        if adult:
            items.append({'key': 'adult', 'label': 'Adult', 'qty': adult, 'price': 500,
                          'seatsPerTicket': 1, 'subtotal': 500 * adult})
        if child:
            items.append({'key': 'child', 'label': 'Child', 'qty': child, 'price': 250,
                          'seatsPerTicket': 1, 'subtotal': 250 * child})
        subtotal = 500 * adult + 250 * child
        order = EventOrder.objects.create(
            event=event, event_date=event_date, user=self.asha,
            event_name=event.name, date=event_date.date, location='Open Grounds',
            start_time=event.start_time, end_time=event.end_time,
            customer_name=customer, phone=phone, items=items,
            seats=adult + child, ticket_count=adult + child,
            subtotal=subtotal, fee=20, amount=subtotal + 20, status=status,
            razorpay_order_id=f'order_{event.slug}_{status}_{pay}_{adult}',
            razorpay_payment_id=pay,
        )
        if age:
            EventOrder.objects.filter(pk=order.pk).update(
                created_at=timezone.now() - datetime.timedelta(minutes=age)
            )
        return order

    def audit(self, count):
        AuditEntry.objects.bulk_create(
            AuditEntry(admin='Admin', action='Did a thing', target=f'T{i}', change='x')
            for i in range(count)
        )
