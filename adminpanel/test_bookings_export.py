"""
GET /api/admin/bookings/export — the Bookings page as an .xlsx.

The rules under test:
  * the export lists exactly the rows the page lists for the same filters
    (type, status, date preset / range on either date basis, search)
  * venue and event rows are never duplicated or mixed up
  * soft-removed rows stay out, as they do on the page
  * the file is a real workbook: frozen, filterable header; dates and
    amounts as typed values; blank cells where a field does not apply
  * only an authenticated super admin can export, and the export is audited
  * nothing about any booking changes
"""
import datetime
import io

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from bookings.models import Booking
from bookings.slots import IST, today_ist
from bookings.views import PENDING_HOLD_MINUTES
from events.models import EventOrder
from events.tests import (
    EventTestCase,
    admin_client,
    auth,
    make_admin,
    make_customer,
    make_event,
    make_vendor,
)

from .exports import COLUMNS, MAX_EXPORT_ROWS
from .models import AuditEntry

URL = '/api/admin/bookings/export'
HEADERS = [c[0] for c in COLUMNS]
OLD = PENDING_HOLD_MINUTES + 5


class BookingsExportTests(EventTestCase):

    def setUp(self):
        super().setUp()
        self.admin = make_admin()
        self.client = admin_client(self.admin)
        self.customer = make_customer()
        self.today = today_ist()
        self.event = make_event(name='Garba Night', days_ahead=(1,))
        self.past_event = make_event(name='Dandiya Past', days_ahead=(-2,))

        V = self.venue
        self.v_paid = V('confirmed', days=2, name='Asha', venue='Grand Hall',
                        pay='pay_v1', method='upi')
        self.v_done = V('confirmed', days=-3, name='Ravi', venue='City Turf',
                        pay='pay_v2', method='card')
        self.v_cancel = V('cancelled', days=1, name='Meena', venue='Grand Hall')
        self.v_pending = V('payment_pending', days=4, name='Kiran', venue='Sky Deck')
        self.v_abandoned = V('payment_pending', days=4, name='Kiran',
                             venue='Sky Deck', age=OLD)
        self.v_refunded = V('refunded', days=-1, name='Asha', venue='City Turf',
                            pay='pay_v3', method='upi')
        self.v_removed = V('removed', days=1, name='Gone', venue='Grand Hall')

        E = self.order
        self.e_paid = E('confirmed', pay='pay_e1', name='Asha')
        self.e_failed = E('failed', name='Ravi')
        self.e_pending = E('payment_pending', name='Kiran')
        self.e_abandoned = E('payment_pending', name='Kiran', age=OLD)
        self.e_cancel = E('cancelled', name='Meena')
        self.e_refund_pending = E('refund_pending', pay='pay_e2', name='Asha')
        self.e_refunded = E('refunded', pay='pay_e3', name='Ravi')
        self.e_done = E('confirmed', pay='pay_e4', name='Ravi', event=self.past_event)
        self.e_removed = E('removed', name='Gone')

    # ---- fixtures ---------------------------------------------------------

    def venue(self, status, *, days, name, venue, pay='', method='online', age=0):
        booking = Booking.objects.create(
            user=self.customer, date=self.today + datetime.timedelta(days=days),
            slots=['19:00 – 20:00', '20:00 – 21:00'], amount=1220, fee=20,
            customer_name=name, venue_name=venue, phone='9100000001',
            status=status, razorpay_payment_id=pay, method=method,
            unit_label='Screen 2', addons=[{'name': 'Cake', 'qty': 1, 'price': 500}],
        )
        if age:
            Booking.objects.filter(pk=booking.pk).update(
                created_at=timezone.now() - datetime.timedelta(minutes=age)
            )
        return booking

    def order(self, status, *, name, pay='', age=0, event=None):
        event = event or self.event
        event_date = event.dates.first()
        order = EventOrder.objects.create(
            event=event, event_date=event_date, user=self.customer,
            event_name=event.name, date=event_date.date, location='Open Grounds',
            start_time=event.start_time, end_time=event.end_time,
            customer_name=name, phone='9100000002',
            items=[{'key': 'adult', 'label': 'Adult', 'qty': 2, 'price': 500,
                    'seatsPerTicket': 1, 'subtotal': 1000},
                   {'key': 'child', 'label': 'Child', 'qty': 1, 'price': 250,
                    'seatsPerTicket': 1, 'subtotal': 250}],
            seats=3, ticket_count=3, subtotal=1250, fee=20, amount=1270,
            status=status, razorpay_order_id=f'order_{name}_{status}_{age}',
            razorpay_payment_id=pay,
        )
        if age:
            EventOrder.objects.filter(pk=order.pk).update(
                created_at=timezone.now() - datetime.timedelta(minutes=age)
            )
        return order

    def set_created(self, obj, days_ago):
        when = timezone.now() - datetime.timedelta(days=days_ago)
        type(obj).objects.filter(pk=obj.pk).update(created_at=when)

    # ---- helpers ----------------------------------------------------------

    def export(self, client=None, **params):
        return (client or self.client).get(URL, params)

    def sheet(self, response):
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        return load_workbook(io.BytesIO(response.content))['Bookings']

    def rows(self, **params):
        ws = self.sheet(self.export(**params))
        values = list(ws.iter_rows(values_only=True))
        self.assertEqual(list(values[0]), HEADERS)
        return [dict(zip(HEADERS, row)) for row in values[1:]]

    def ids(self, **params):
        return [r['Record ID'] for r in self.rows(**params)]

    def assertExports(self, expected, **params):
        got = self.ids(**params)
        self.assertEqual(len(got), len(set(got)), 'duplicate rows')
        self.assertEqual(set(got), {o.id for o in expected}, params)

    ALL_LIVE = (
        'v_paid', 'v_done', 'v_cancel', 'v_pending', 'v_abandoned', 'v_refunded',
        'e_paid', 'e_failed', 'e_pending', 'e_abandoned', 'e_cancel',
        'e_refund_pending', 'e_refunded', 'e_done',
    )

    def objs(self, *names):
        return [getattr(self, n) for n in names]

    # ---- 1-3: types -------------------------------------------------------

    def test_exports_all_bookings_without_removed_rows(self):
        self.assertExports(self.objs(*self.ALL_LIVE))
        ids = self.ids()
        self.assertNotIn(self.v_removed.id, ids)
        self.assertNotIn(self.e_removed.id, ids)

    def test_exports_only_event_bookings(self):
        rows = self.rows(type='event')
        self.assertEqual({r['Booking Type'] for r in rows}, {'Event ticket'})
        self.assertExports(self.objs(*[n for n in self.ALL_LIVE if n[0] == 'e']),
                           type='event')

    def test_exports_only_venue_bookings(self):
        rows = self.rows(type='venue')
        self.assertEqual({r['Booking Type'] for r in rows}, {'Venue booking'})
        self.assertExports(self.objs(*[n for n in self.ALL_LIVE if n[0] == 'v']),
                           type='venue')

    # ---- 4: statuses ------------------------------------------------------

    def test_each_status_exports_exactly_its_rows(self):
        cases = {
            'confirmed': ('v_paid', 'e_paid'),
            'completed': ('v_done', 'e_done'),
            'payment_pending': ('v_pending', 'e_pending'),
            'expired': ('v_abandoned', 'e_abandoned'),
            'failed': ('e_failed',),
            'cancelled': ('v_cancel', 'e_cancel'),
            'refund_pending': ('e_refund_pending',),
            'refunded': ('v_refunded', 'e_refunded'),
        }
        for status, names in cases.items():
            with self.subTest(status=status):
                self.assertExports(self.objs(*names), status=status)

    def test_status_columns_keep_the_stored_value(self):
        rows = {r['Record ID']: r for r in self.rows()}
        self.assertEqual(rows[self.v_done.id]['Payment Status'], 'confirmed')
        self.assertEqual(rows[self.v_done.id]['Booking Status'], 'Completed')
        self.assertEqual(rows[self.v_abandoned.id]['Payment Status'], 'payment_pending')
        self.assertEqual(rows[self.v_abandoned.id]['Booking Status'], 'Abandoned checkout')
        self.assertEqual(rows[self.e_paid.id]['Booking Status'], 'Paid')
        self.assertEqual(rows[self.e_failed.id]['Payment Status'], 'failed')
        self.assertEqual(rows[self.e_failed.id]['Booking Status'], 'Payment failed')
        self.assertEqual(rows[self.e_abandoned.id]['Booking Status'], 'Abandoned checkout')

    def test_failed_venue_payment_exports_as_payment_failed(self):
        failed = self.venue('failed', days=3, name='Nina', venue='Grand Hall')
        rows = self.rows(type='venue', status='failed')
        self.assertEqual([r['Record ID'] for r in rows], [failed.id])
        self.assertEqual(rows[0]['Payment Status'], 'failed')
        self.assertEqual(rows[0]['Booking Status'], 'Payment failed')
        # Never counted as paid.
        self.assertNotIn(failed.id, self.ids(status='confirmed'))

    def test_event_rows_carry_their_short_ids(self):
        rows = {r['Record ID']: r for r in self.rows(type='event')}
        for order in self.objs('e_paid', 'e_failed', 'e_abandoned'):
            self.assertRegex(order.id, r'^[A-Z0-9]{6}$')
            self.assertEqual(rows[order.id]['Booking ID'], f'#{order.id}')
        self.assertExports(self.objs('e_paid'), q=f'#{self.e_paid.id}'.lower())

    # ---- 5: dates ---------------------------------------------------------

    def test_event_date_range_selects_the_event_day(self):
        day = self.event.dates.first().date.isoformat()
        self.assertExports(
            self.objs('e_paid', 'e_failed', 'e_pending', 'e_abandoned',
                      'e_cancel', 'e_refund_pending', 'e_refunded'),
            type='event', date='range', **{'from': day, 'to': day},
        )

    def test_venue_scheduled_date_range(self):
        start = (self.today + datetime.timedelta(days=1)).isoformat()
        end = (self.today + datetime.timedelta(days=2)).isoformat()
        self.assertExports(self.objs('v_paid', 'v_cancel'),
                           type='venue', date='range', **{'from': start, 'to': end})

    def test_open_ended_range(self):
        start = (self.today + datetime.timedelta(days=3)).isoformat()
        self.assertExports(self.objs('v_pending', 'v_abandoned'),
                           type='venue', date='range', **{'from': start})

    def test_presets_split_on_the_scheduled_date(self):
        self.assertExports(self.objs('v_done', 'v_refunded', 'e_done'), date='past')
        self.assertExports([], date='today')
        upcoming = [n for n in self.ALL_LIVE if n not in ('v_done', 'v_refunded', 'e_done')]
        self.assertExports(self.objs(*upcoming), date='upcoming')

    def test_created_date_basis_reads_the_booking_day(self):
        self.set_created(self.v_paid, 10)
        self.set_created(self.e_failed, 10)
        day = (timezone.now() - datetime.timedelta(days=10)).astimezone(IST).date().isoformat()
        self.assertExports(self.objs('v_paid', 'e_failed'),
                           date='range', basis='created', **{'from': day, 'to': day})
        # The same day as a scheduled date matches nothing.
        self.assertExports([], date='range', basis='scheduled', **{'from': day, 'to': day})

    def test_created_basis_uses_ist_calendar_days(self):
        # 00:30 IST on day D is 19:00 UTC on day D-1: the page shows it on
        # day D, so the filter must cut days in IST, not UTC.
        day = self.today - datetime.timedelta(days=5)
        early = datetime.datetime.combine(day, datetime.time(0, 30), tzinfo=IST)
        Booking.objects.filter(pk=self.v_cancel.pk).update(created_at=early)
        iso = day.isoformat()
        self.assertExports(self.objs('v_cancel'),
                           date='range', basis='created', **{'from': iso, 'to': iso})

    # ---- 6: combined ------------------------------------------------------

    def test_filters_combine(self):
        day = self.event.dates.first().date.isoformat()
        self.assertExports(self.objs('e_failed'), type='event', status='failed',
                           date='range', **{'from': day, 'to': day})
        self.assertExports(self.objs('v_paid'), type='venue', status='confirmed',
                           date='upcoming', q='grand')
        self.assertExports([], type='venue', status='failed')

    # ---- 7: search --------------------------------------------------------

    def test_search_by_booking_id_and_names(self):
        self.assertExports(self.objs('v_paid'), q=self.v_paid.id)
        ref = '#' + self.v_paid.id     # the shared 6-character booking ID
        self.assertExports(self.objs('v_paid'), q=ref.lower())
        self.assertExports(self.objs('e_paid'), q=self.e_paid.id.upper())
        self.assertExports(self.objs('v_pending', 'v_abandoned', 'e_pending',
                                     'e_abandoned'), q='kiran')
        self.assertExports(self.objs('v_done', 'v_refunded'), q='city turf')
        self.assertExports(self.objs('e_done'), q='dandiya past')
        self.assertExports(self.objs('e_refund_pending'), q='pay_e2')

    def test_older_long_event_ids_export_and_search_as_before(self):
        legacy = self.order('confirmed', pay='pay_old', name='Old Timer')
        EventOrder.objects.filter(pk=legacy.pk).update(id='eo_100a524e06f6')
        rows = self.rows(q='eo_100a524e06f6')
        self.assertEqual([r['Record ID'] for r in rows], ['eo_100a524e06f6'])
        self.assertEqual(rows[0]['Booking ID'], '#4E06F6')
        self.assertEqual([r['Record ID'] for r in self.rows(q='#4e06f6')], ['eo_100a524e06f6'])

    # ---- 8: access --------------------------------------------------------

    def test_only_a_super_admin_can_export(self):
        for client in (APIClient(), auth(make_customer('9100000055')),
                       auth(make_vendor()),
                       admin_client(make_customer('9100000044'))):
            response = self.export(client)
            self.assertIn(response.status_code, (401, 403))
            self.assertNotIn('spreadsheet', response.get('Content-Type', ''))
        self.assertFalse(AuditEntry.objects.filter(action='Exported bookings').exists())

    def test_bad_filters_are_refused(self):
        for params in ({'type': 'hotel'}, {'status': 'paid_maybe'},
                       {'date': 'yesterday'}, {'basis': 'paid'},
                       {'date': 'range', 'from': '2026-13-01'},
                       {'date': 'range', 'from': '2026-10-10', 'to': '2026-10-01'}):
            with self.subTest(params=params):
                self.assertEqual(self.export(**params).status_code, 400)

    def test_export_is_audited(self):
        self.export(status='failed')
        entry = AuditEntry.objects.get(action='Exported bookings')
        self.assertEqual(entry.target, '1 bookings')
        self.assertIn('status=failed', entry.change)

    def test_oversized_export_is_refused(self):
        from . import exports
        original = exports.MAX_EXPORT_ROWS
        exports.MAX_EXPORT_ROWS = 3
        try:
            response = self.export()
        finally:
            exports.MAX_EXPORT_ROWS = original
        self.assertEqual(response.status_code, 400)
        self.assertIn('Narrow the filters', response.data['detail'])
        self.assertEqual(MAX_EXPORT_ROWS, original)

    # ---- 9: agrees with what the page lists -------------------------------

    def test_export_matches_the_page_rows_for_every_type_and_status(self):
        """The page filters the bootstrap rows by statusKey; the export must
        list the same ids for the same filter."""
        data = self.client.get('/api/admin/bootstrap').data
        page = [('venue', b['id'], 'expired' if b['status'] == 'payment_pending'
                 and b['holdExpired'] else b['status']) for b in data['bookings']]
        page += [('event', o['id'], o['status']) for o in data['eventOrders']]
        statuses = {'All'} | {key for _k, _i, key in page}
        for kind in ('All', 'venue', 'event'):
            for status in statuses:
                with self.subTest(kind=kind, status=status):
                    expected = sorted(
                        i for k, i, key in page
                        if kind in ('All', k) and status in ('All', key)
                    )
                    self.assertEqual(sorted(self.ids(type=kind, status=status)), expected)

    # ---- 10: the file -----------------------------------------------------

    def test_file_is_a_valid_formatted_workbook(self):
        response = self.export()
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertEqual(
            response['Content-Disposition'],
            f'attachment; filename="bookmyvenues_bookings_{self.today.isoformat()}.xlsx"',
        )
        self.assertEqual(response['Cache-Control'], 'no-store')
        self.assertTrue(response.content.startswith(b'PK'))  # a zip container
        ws = self.sheet(response)
        self.assertEqual(ws.freeze_panes, 'A2')
        self.assertEqual(ws.auto_filter.ref, f'A1:O{len(self.ALL_LIVE) + 1}')
        self.assertTrue(ws['A1'].font.bold)
        self.assertGreaterEqual(ws.column_dimensions['D'].width, 30)
        ids = [c.value for c in ws['B'][1:]]
        self.assertEqual(len(ids), len(set(ids)))

    def test_cells_are_typed_and_blank_when_not_applicable(self):
        ws = self.sheet(self.export())
        by_id = {row[1].value: row for row in ws.iter_rows(min_row=2)}
        col = {h: i for i, h in enumerate(HEADERS)}

        venue = by_id[self.v_paid.id]
        self.assertEqual(venue[col['Event / Booking Date']].value.date(),
                         self.v_paid.date)
        self.assertEqual(venue[col['Event / Booking Date']].number_format, 'dd-mmm-yyyy')
        self.assertEqual(venue[col['Booking Amount']].value, 1220)
        self.assertIn('#,##0', venue[col['Booking Amount']].number_format)
        self.assertIsNone(venue[col['Quantity']].value)       # not a ticket
        self.assertEqual(venue[col['Time Slot']].value, '19:00 – 20:00, 20:00 – 21:00')
        self.assertIn('Screen 2', venue[col['Ticket / Booking Details']].value)
        self.assertIn('Cake', venue[col['Ticket / Booking Details']].value)
        self.assertEqual(venue[col['Payment Method']].value, 'UPI')
        self.assertEqual(venue[col['Customer Phone']].value, '9100000001')
        created = venue[col['Booking Created At']].value
        self.assertIsInstance(created, datetime.datetime)
        expected = self.v_paid.created_at.astimezone(IST).replace(tzinfo=None)
        self.assertLess(abs((created - expected).total_seconds()), 1)

        event = by_id[self.e_paid.id]
        self.assertEqual(event[col['Venue / Event Name']].value, 'Garba Night')
        self.assertEqual(event[col['Quantity']].value, 3)
        self.assertEqual(event[col['Time Slot']].value, '18:00 – 22:00')
        self.assertEqual(event[col['Ticket / Booking Details']].value, '2× Adult, 1× Child')
        self.assertEqual(event[col['Booking Amount']].value, 1270)
        self.assertEqual(event[col['Payment Method']].value, 'Online')
        # No payment recorded -> no method, rather than a guessed one.
        self.assertIsNone(by_id[self.e_failed.id][col['Payment Method']].value)

    def test_empty_export_is_still_a_valid_workbook(self):
        ws = self.sheet(self.export(q='no-such-booking'))
        self.assertEqual([c.value for c in ws[1]], HEADERS)
        self.assertEqual(ws.max_row, 1)

    def test_rows_are_newest_first(self):
        self.set_created(self.v_paid, 3)
        self.set_created(self.e_paid, 1)
        ids = self.ids(status='confirmed')
        self.assertEqual(ids, [self.e_paid.id, self.v_paid.id])

    def test_export_changes_nothing(self):
        before = {
            'b': sorted(Booking.objects.values_list('id', 'status', 'amount')),
            'e': sorted(EventOrder.objects.values_list('id', 'status', 'amount')),
        }
        self.export()
        self.export(type='event', status='expired')
        after = {
            'b': sorted(Booking.objects.values_list('id', 'status', 'amount')),
            'e': sorted(EventOrder.objects.values_list('id', 'status', 'amount')),
        }
        self.assertEqual(before, after)

    def test_route_does_not_shadow_booking_updates(self):
        self.assertEqual(reverse('admin-bookings-export'), URL)
        response = self.client.patch(
            reverse('admin-booking', args=[self.v_cancel.id]), {}, format='json'
        )
        self.assertNotEqual(response.status_code, 404)
