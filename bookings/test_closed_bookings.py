"""
Paid bookings that did not go ahead — refunded, or cancelled after a payment
was captured — are listed for the vendor (dashboard `closedBookings`) and
exported with the "Cancelled & refunded" filter. Abandoned, never-paid
checkouts are not bookings and stay out.
"""
import datetime
import io

from openpyxl import load_workbook

from .models import Booking
from .slots import today_ist
from .tests import BookingTestBase


class ClosedBookingsTests(BookingTestBase):
    def setUp(self):
        super().setUp()
        day = today_ist() - datetime.timedelta(days=2)
        base = dict(listing=self.listing, user=self.customer, date=day, slots=['10:00 – 11:00'],
                    amount=620, venue_name='Grand Palace Hall', customer_name='Asha')
        self.refunded = Booking.objects.create(**base, status='refunded', razorpay_payment_id='pay_r',
                                               refund_amount=620)
        self.paid_cancel = Booking.objects.create(**base, status='cancelled', razorpay_payment_id='pay_c')
        self.abandoned = Booking.objects.create(**base, status='cancelled')
        self.kept = Booking.objects.create(**base, status='confirmed', razorpay_payment_id='pay_k')
        self.client.force_authenticate(user=self.vendor)

    def test_dashboard_lists_paid_bookings_that_did_not_go_ahead(self):
        data = self.client.get('/api/vendors/me/dashboard').data
        closed = {b['id'] for b in data['closedBookings']}
        self.assertEqual(closed, {self.refunded.id, self.paid_cancel.id})
        standing = {b['id'] for b in data['allBookings']}
        self.assertEqual(standing, {self.kept.id})      # unchanged rule

    def test_export_with_the_cancelled_filter(self):
        response = self.client.get('/api/vendors/me/exports/bookings?status=cancelled')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(b''.join(response.streaming_content)
                                         if hasattr(response, 'streaming_content') else response.content)).active
        ids = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True) if row and row[0]}
        self.assertTrue({self.refunded.id, self.paid_cancel.id} <= ids)
        self.assertNotIn(self.abandoned.id, ids)
        self.assertNotIn(self.kept.id, ids)
